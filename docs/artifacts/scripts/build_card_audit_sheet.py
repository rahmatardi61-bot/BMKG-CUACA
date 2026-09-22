#!/usr/bin/env python3
# ═══════════════════════════════════════════════════════════════════════════
# Build sheet "Audit Card + Screenshot" di docs/bmkg-api-mapping.xlsx
# ---------------------------------------------------------------------------
# Sumber data:
#   docs/screenshots/cards/*.png   — screenshot tiap component/card (Playwright)
#   docs/screenshots/cards/dump.json — hasil window.__apiBox (card→API, URL, sample)
#
# Regenerate:
#   1) npm run build && node server.mjs            (port 8080, data live)
#   2) node docs/scrapping_cuaca-bmkg-go-id/capture_card_screenshots.js
#   3) python3 docs/artifacts/scripts/build_card_audit_sheet.py
# ═══════════════════════════════════════════════════════════════════════════
import json
import os
from PIL import Image
import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont

BASE = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..', '..'))  # repo root (docs/artifacts/scripts → root)
CARD_DIR = os.path.join(BASE, 'docs', 'screenshots', 'cards')
XLSX = os.path.join(BASE, 'docs', 'bmkg-api-mapping.xlsx')
SHEET = 'Audit Card + Screenshot'
G_BUDGET = 32000  # budget isi kolom G per kartu (batas sel Excel 32.767; sisanya di sheet Sample Full)

# Field response yang dipakai UI + penjelasannya (sumber kebenaran:
# src/services/bmkg/adapters.ts, openData.ts, dan komponen terkait)
API_USAGE = {
  'df-forecast': ('data[0].cuaca[][] → t, hu, weather_desc, tp, ws, tcc, local_datetime',
    'Suhu & kelembapan; weather_desc → ikon lucide (descToIcon); hujan mm → % (tp×20, aproksimasi); local_datetime → urutan jam & strip 7 hari (slot pertama tiap grup). API tidak punya tmin/tmax/UV → diestimasi.'),
  'df-adm': ('provinsi, kotkab, kecamatan, desa',
    'Resolusi koordinat → wilayah administratif; label lokasi kartu.'),
  'df-amandemen': ('data[] (item cuaca amandemen terbaru)',
    'Pembaruan prakiraan (amandemen) → menyegarkan nilai/badge peringatan.'),
  'presentwx': ('data.t, hu, weather_desc, ws, tcc, vs, local_datetime',
    'Kondisi terobservasi saat ini: suhu, status→ikon, "terasa seperti" = t+(hu-50)/10, UV estimasi dari tcc, jarak pandang vs m→km.'),
  'sunset': ('sunrise, sunset (format 12-jam)',
    'Dikonversi ke HH:mm → info matahari terbit/terbenam di CurrentWeather.'),
  'public-warning': ('data.today / data.tomorrow (warning_level, impact/description)',
    'Peringatan hari ini/besok → severity Awas/Siaga/Waspada → AlertsPanel; kosong = tanpa peringatan.'),
  'public-warning-cyclone': ('data siklon terdampak',
    'Peringatan siklon untuk wilayah — hanya muncul saat ada siklon aktif.'),
  'nowcast': ('XML RSS: <item> per provinsi (title, description, pubDate)',
    'Peringatan dini 3-6 jam resmi; difilter per provinsi; menang atas warning internal df.'),
  'blog-wp': ('[].title.rendered, excerpt.rendered, link, _embedded (featured media)',
    'Artikel berita WP → NewsSection; fallback mock saat upstream 502 (sering).'),
  'public-video': ('[].judul, video_url',
    'Video berita terbaru (per hashtag) → tab video NewsSection.'),
  'tews': ('features[].properties (magnitude, kedalaman, wilayah, waktu)',
    'Autogempa terkini + gempa dirasakan → EarthquakePanel.'),
  'dwt': ('rows [kota, kecamatan/stasiun, lat, lng, id, waktu, slots cuaca[]]',
    'Cuaca kecamatan (jalan raya) & stasiun (kereta) terdekat ≤2° → kartu transportasi.'),
  'maritim-nearest': ('code / nama wilayah laut terdekat',
    'Anchor koordinat → wilayah perairan/pelabuhan resmi terdekat.'),
  'maritim-perairan': ('data[] → weather_desc, wave_desc, wind_speed_min/max, warning_desc, time_desc',
    'Slot perairan resmi (Hari ini…H+3) → kartu maritim & aktivitas pelayaran; currentSlot = slot pertama yang masih berlaku.'),
  'maritim-overview': ('{code: issued, today, tomorrow, wave_desc}',
    'Ringkasan 232 wilayah → warna peta laut + legend gelombang.'),
  'maritim-wilayah': ('features[].geometry (polygon 232 wilayah)',
    'Polygon GeoJSON → layer batas wilayah perairan di peta.'),
  'maritim-pelabuhan-list': ('files[].coor, file',
    'Koordinat 294 pelabuhan → pilih terdekat → fetch detail.'),
  'maritim-pelabuhan': ('detail pelabuhan + data pasut',
    'Pasut (UTC → zona waktu pelabuhan) → PortTideCard; pasut null → blok disembunyikan.'),
  'official-forecast': ('data[].cuaca[][] → t, hu, weather_desc, tp',
    'Fallback prakiraan resmi (open-data) bila df kosong — bentuk item identik df.'),
  'overpass': ('elements[].lat/lon/tags',
    'POI OpenStreetMap sekitar pengguna → AroundActivity.'),
  'nominatim': ('[].lat/lon/display_name',
    'Geocode nama tempat → koordinat.'),
  'osrm': ('routes[].geometry, duration, distance',
    'Rute antar titik → panel aktivitas darat.'),
  'satellite-modelrun': ('himawari9[] (baserun ISO UTC, terbaru index 0)',
    'Satellite Tile API — daftar 18 frame @10 menit → 18 TileLayer pra-muat; play = toggle opacity (0 request saat animasi). CORS *, tanpa proxy.'),
  'satelit': ('(PNG biner) inderaja.bmkg.go.id H08_*.png',
    'Citra Himawari statis → FALLBACK peta satelit bila Tile API gagal (background <img>; terekam sebagai request gambar biner).'),
  'tcwc-cyclone': ('data siklon tropis aktif',
    'Posisi & intensitas siklon → panel siklon.'),
  'local-json': ('file statis /data|sectors|ports.json',
    'Snapshot lokal (layer angin, pelabuhan, landmark) → bagian kartu yang mock.'),
}

JUDUL = {
    'hero': 'Hero Weather Card — Cuaca Hari Ini',
    'current-weather': 'Kondisi Saat Ini (rincian)',
    'forecast-panel': 'Panel Prakiraan (Per Jam / Mingguan)',
    'alerts': 'Peringatan Dini (Alerts)',
    'news': 'Berita & Video (News Section)',
    'earthquake': 'Gempa Bumi (TEWS)',
    'satellite': 'Citra Satelit Himawari-9',
    'major-cities': 'Carousel Kota Besar (statis)',
    'transport': 'Transportasi (Jalan / Kereta / Maritim)',
    'weather-activity': 'Panduan Aktivitas & Analisis Cuaca',
    'marine-map': 'Peta Laut (Marine Map)',
    'port-tide': 'Pelabuhan & Pasut',
    'around-activity': 'Aktivitas Sekitar (POI)',
    'land-based': 'Aktivitas Darat (Drawer)',
    'maritime-advisor': 'Konsultasi Pelayaran (Drawer)',
    'aviation-advisor': 'Konsultasi Penerbangan (Drawer)',
    'location-search': 'Pencarian Lokasi',
}
STATUS = {
    'live': ('✅ LIVE', 'C6EFCE'),
    'mixed': ('⚠️ CAMPURAN', 'FFEB9C'),
    'mock': ('🔌 MOCK', 'E7E6E6'),
}

dump = json.load(open(os.path.join(CARD_DIR, 'dump.json'), encoding='utf-8'))
cards, defs, calls = dump['cards'], dump['defs'], dump['calls']

# ── siapkan JPEG ter-resize untuk embed (PNG 2x terlalu berat) ──────────────
tmp = os.path.join(CARD_DIR, '.sheet')
os.makedirs(tmp, exist_ok=True)
DISP_W = 460  # px tampilan (kecil, cukup untuk mengenali komponen)
for card_id in cards:
    src = os.path.join(CARD_DIR, f'{card_id}.png')
    if not os.path.exists(src):
        continue
    im = Image.open(src).convert('RGB')
    w, h = im.size
    nh = round(h * DISP_W / w)
    im.resize((DISP_W, nh), Image.LANCZOS).save(os.path.join(tmp, f'{card_id}.jpg'), quality=82)

# ── rakit sheet ──────────────────────────────────────────────────────────────
wb = openpyxl.load_workbook(XLSX)
if SHEET in wb.sheetnames:
    del wb[SHEET]
idx = wb.sheetnames.index('Audit Content Card') + 1
ws = wb.create_sheet(SHEET, index=idx)

thin = Border(*[Side(style='thin', color='D0D5DD')] * 4)
wrap = Alignment(wrap_text=True, vertical='top')
head_fill = PatternFill('solid', fgColor='1F3864')
head_font = Font(bold=True, color='FFFFFF', size=10)

ws.merge_cells('A1:I1')
ws['A1'] = 'AUDIT CONTENT CARD + SCREENSHOT — komponen, API yang dipakai, query, sample response'
ws['A1'].font = Font(bold=True, size=13)
ws.merge_cells('A2:I2')
ws['A2'] = ('Screenshot diambil dari deployment prod (server.mjs + dist, data live) dengan API BOX MARKER aktif '
            '— border merah = live, amber = campuran, abu = mock (lihat kolom Status). '
    'Kolom "Sample Response" = respons asli utuh selama muat dalam satu sel (batas Excel 32.767 char/kartu) — sisanya & salinan lengkap per API: sheet "Sample Full". '
    'Kolom "Field dipakai UI" = pemetaan response → nilai yang dirender (sumber: src/services/bmkg/adapters.ts). '
    '*(proxy) merah = API tidak bisa diakses langsung dari browser (CORS / header khusus) — harus lewat proxy/alias untuk ambil response-nya. ''live = respon asli API · estimasi = nilai turunan/konversi dari API · mock = data statis (endpoint belum tersedia). '
            'Sumber screenshot: docs/screenshots/cards/ · skrip: docs/scrapping_cuaca-bmkg-go-id/capture_card_screenshots.js')
ws['A2'].font = Font(size=9, italic=True, color='555555')
ws.merge_cells('A3:I3')
ws['A3'] = 'Status: ✅ LIVE = nilai kartu dari API live  |  ⚠️ CAMPURAN = sebagian live/sebagian mock  |  🔌 MOCK = statis/estimasi'
ws['A3'].font = Font(size=9, bold=True, color='555555')

headers = ['No', 'Komponen', 'Screenshot', 'Status', 'API dipakai', 'Query / URL terakhir', 'Sample Response (per API, utuh selama muat di sel)', 'Field dipakai UI + penjelasan', 'Catatan']
for c, h in enumerate(headers, 1):
    cell = ws.cell(row=4, column=c, value=h)
    cell.fill, cell.font, cell.border = head_fill, head_font, thin
    cell.alignment = Alignment(horizontal='center', vertical='center')

row = 5
for no, (card_id, m) in enumerate(cards.items(), 1):
    label, fill = STATUS[m['kind']]
    judul = JUDUL.get(card_id, card_id)

    ws.cell(row=row, column=1, value=no)
    ws.cell(row=row, column=2, value=f'{judul}\n({card_id})')

    img_path = os.path.join(tmp, f'{card_id}.jpg')
    img_h = 0
    if os.path.exists(img_path):
        im = XLImage(img_path)
        img_h = im.height
        im.anchor = f'C{row}'
        ws.add_image(im)
    else:
        ws.cell(row=row, column=3, value='(screenshot tidak tersedia)')

    sc = ws.cell(row=row, column=4, value=label)
    sc.fill = PatternFill('solid', fgColor=fill)
    sc.alignment = Alignment(vertical='top')

    # E: API dipakai — F: query/url — G: sample PER API (utuh selama muat di sel) — H: field
    e_lines, f_lines, g_blocks, h_lines = [], [], [], []
    names, raws = {}, {}
    for api_id in m['apis']:
        name = defs.get(api_id, {}).get('name', api_id)
        names[api_id] = name
        e_lines.append(f'▸ {name}')
        call = (calls.get(api_id) or [{}])[-1]
        if call:
            f_lines.append((f'▸ [{call.get("status", "?")}] {call.get("url", "?")}', call.get('proxied', False)))
            raws[api_id] = ' '.join(call.get('sample', '').split())  # respons asli, dirapikan 1 baris
        else:
            f_lines.append((f'▸ {name} — (belum terekam saat capture)', False))
            raws[api_id] = ''
        fields, penjelasan = API_USAGE.get(api_id, ('(belum terpetakan)', ''))
        h_lines.append(f'▸ {name}\n  Dipakai: {fields}\n  → {penjelasan}')

    # G diisi respons UTUH selama total muat dalam satu sel (batas Excel 32.767 char):
    # pass 1 = preview 600 char per API, pass 2 = sisa budget mengalir ke API berurutan
    shown = {aid: min(600, len(raws[aid])) for aid in raws}
    used = sum(len(names[aid]) + shown[aid] + 8 for aid in raws)
    room = G_BUDGET - used
    for aid in raws:
        if room <= 0:
            break
        take = min(len(raws[aid]) - shown[aid], room)
        shown[aid] += take
        room -= take
    for aid in raws:
        name, raw = names[aid], raws[aid]
        if not raw:
            g_blocks.append(f'▸ {name}\n(tidak ada response live)')
            continue
        pot = raw[:shown[aid]]
        sisa = len(raw) - shown[aid]
        g_blocks.append(f'▸ {name}\n{pot}' + (f' … [+{sisa} char — utuh di sheet "Sample Full"]' if sisa > 0 else ''))
    ws.cell(row=row, column=5, value='\n'.join(e_lines) if e_lines else '— (tanpa API live)')
    # kolom F: URL asli upstream; ' *(proxy)' merah tebal = API butuh proxy/alias
    # (CORS / header khusus) kalau mau ambil response-nya
    if any(pr for _, pr in f_lines):
        fnt = InlineFont(rFont='Courier New', sz=8)
        red = InlineFont(rFont='Courier New', sz=8, b=True, color='FFCC0000')
        rt = CellRichText()
        for i, (line, proxied) in enumerate(f_lines):
            # newline digabung ke awal teks baris berikutnya dalam run yang sama —
            # run berisi '\n' saja di-drop Excel (baris menyatu); pola file Excel asli
            # = '\n' menempel konten dalam <t xml:space="preserve">
            text = line if i == 0 else '\n' + line
            rt.append(TextBlock(fnt, text))
            if proxied:
                rt.append(TextBlock(red, ' *(proxy)'))
        ws.cell(row=row, column=6, value=rt)
    else:
        ws.cell(row=row, column=6, value='\n'.join(l for l, _ in f_lines) if f_lines else '—')
    ws.cell(row=row, column=7, value='\n\n'.join(g_blocks) if g_blocks else '— (tidak ada response live)')
    ws.cell(row=row, column=8, value='\n'.join(h_lines) if h_lines else '— (kartu tanpa API live — 100% statis)')

    note = m.get('note', '')
    parts = m.get('parts') or []
    i_lines = [note] if note else []
    i_lines += [f"▸ {p['label']} — " + {'live': 'data asli dari API', 'est': 'estimasi/konversi dari API', 'mock': 'data statis (belum ada endpoint)'}[p['type']] for p in parts]
    ws.cell(row=row, column=9, value='\n'.join(i_lines) or '—')

    # format baris
    for c in range(1, 10):
        cell = ws.cell(row=row, column=c)
        cell.border = thin
        if c in (2, 5, 6, 7, 8, 9):
            cell.alignment = wrap
        if c in (1, 2, 5, 8):
            cell.font = Font(size=9)
        if c in (6, 7):
            cell.font = Font(size=8, name='Courier New')
    # tinggi baris = konten terpanjang (sample per-API / field+penjelasan / catatan); batas Excel 409,5pt
    def est_h_of(col, chars):
        txt = str(ws.cell(row=row, column=col).value or '')
        return sum(max(1, -(-len(seg) // chars)) for seg in txt.split('\n')) * 10.5 + 6
    est_h = max(est_h_of(6, 52), est_h_of(7, 92), est_h_of(8, 56), est_h_of(9, 38))
    ws.row_dimensions[row].height = min(409, max(90, est_h, img_h * 0.75 + 8)) if img_h else min(409, max(90, est_h))
    row += 1

for col, w in {'A': 5, 'B': 28, 'C': 70, 'D': 14, 'E': 30, 'F': 48, 'G': 92, 'H': 56, 'I': 38}.items():
    ws.column_dimensions[col].width = w
ws.freeze_panes = 'A5'

# ══ sheet 'Sample Full': response mentah UTUH — tanpa dipotong ══
# Batas fisik sel Excel = 32.767 char → respons panjang dipecah per 30.000 char
# (kolom '#' = potongan ke-N; gabungkan berurutan = respons asli utuh).
sf = wb.create_sheet('Sample Full')
sf['A1'] = ('SAMPLE RESPONSE LENGKAP — respons asli utuh dari hasil fetch. '
            'Sel Excel maks 32.767 char → respons panjang dipecah per 30.000 char '
            '(kolom # = urutan potongan). Sample penuh juga tersedia di docs/screenshots/cards/dump.json')
sf['A1'].font = Font(bold=True, size=11)
sf.merge_cells('A1:F1')
for c, (h, w) in enumerate(zip(['Card', 'API', '#', 'Status', 'URL', 'Response (utuh, potongan 30rb char)'],
                               [24, 30, 5, 8, 46, 180]), start=1):
    cell = sf.cell(row=2, column=c, value=h)
    cell.font = Font(bold=True, size=9)
    cell.border = thin
    cell.alignment = Alignment(vertical='center')
    sf.column_dimensions[chr(64 + c)].width = w

card_of = {}
for cid, cd in dump.get('cards', {}).items():
    for aid in cd.get('apis', []):
        card_of.setdefault(aid, cid)

r = 3
for api_id, clist in (dump.get('calls') or {}).items():
    api_name = defs.get(api_id, {}).get('name', api_id)
    for call in clist:
        sample = call.get('sample', '')
        if not sample:
            continue
        chunks = [sample[i:i + 30000] for i in range(0, len(sample), 30000)]
        for ci, ch in enumerate(chunks, start=1):
            vals = [card_of.get(api_id, '—'), api_name, ci, call.get('status', '?'),
                    call.get('url', '?'), ch]
            for c, v in enumerate(vals, start=1):
                cell = sf.cell(row=r, column=c, value=v)
                cell.border = thin
                cell.font = Font(size=8, name='Courier New') if c == 6 else Font(size=9)
            sf.row_dimensions[r].height = 12
            r += 1
sf.freeze_panes = 'A3'

wb.save(XLSX)
print(f"OK — sheet '{SHEET}': {len(cards)} card, gambar {sum(1 for c in cards if os.path.exists(os.path.join(tmp, f'{c}.jpg')))}")
