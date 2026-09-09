#!/usr/bin/env python3
"""Generate api-mapping.xlsx — komparasi API baseline (cuaca.bmkg.go.id) vs project redesign.
Data satu sumber: ROWS di bawah. Re-run: python3 generate_xlsx.py"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# status: BELUM (belum diimplementasikan) | BEDA-SUMBER (ada tapi pakai sumber lain/mock) | REDESIGN-ONLY (sumber yang tidak dipakai original)
ROWS = [
    # (fam, endpoint, method, params, auth, dipakai_original, komponen_redesign, status, sumber_pengganti_redesign, catatan)
    ("C. DF Forecast", "/api/df/v1/forecast/coord", "GET", "lat, lon", "Referer+Origin",
     "Dashboard: cuaca saat ini, prakiraan 10 hari, tab Ringkasan/Per Jam; /perbandingan",
     "CurrentWeather.vue, ForecastPanel.vue, ForecastTemperatureChart.vue, ForecastPrecipitationChart.vue, ForecastAirQualityChart.vue",
     "BELUM", "src/data/mockData.ts (weatherDataMap, hourlyForecastsMap, generateMockWeatherForCity)",
     "Respons: lokasi adm1-adm4 + cuaca per jam & 3-jam hingga +9 hari (t, tcc, tp, weather, wd, ws, hu, vs, image icon). Endpoint paling kritis untuk diimplementasikan."),
    ("C. DF Forecast", "/api/df/v1/adm/coord", "GET", "lat, lon", "Referer+Origin",
     "Resolve lokasi admin (desa/kec/kotkab/prov) dari koordinat",
     "App.vue (geolocation handler)", "BEDA-SUMBER", "Nominatim reverse geocoding (openstreetmap.org)",
     "Struktur respons original: adm1..adm4 + provinsi/kotkab/kecamatan/desa + timezone. Nominatim butuh mapping manual field-nya."),
    ("C. DF Forecast", "/api/df/v1/amandemen/coord", "GET", "lat, lon", "Referer+Origin",
     "Amandemen prakiraan (revisi data)", "-", "BELUM", "-", "Fitur belum ada sama sekali di redesign."),
    ("D. Present Weather", "/api/presentwx/coord", "GET", "lat, lon", "Terbuka",
     "Kartu 'Cuaca Saat Ini' (observasi aktual)", "CurrentWeather.vue", "BELUM",
     "src/data/mockData.ts (weatherDataMap)", "Respons: {status, data:{lokasi, cuaca}}."),
    ("B. Public API", "/api/public/weather/warning", "GET", "lat, long", "x-public-token",
     "Panel peringatan cuaca wilayah (today/tomorrow)", "AlertsPanel.vue", "BELUM",
     "src/data/mockData.ts (warningAlertsMap)", "Respons: {data:{today:{...},tomorrow:{...}}}."),
    ("B. Public API", "/api/public/weather/warning/cyclone", "GET", "lat, long", "x-public-token",
     "Status siklon lokal", "-", "BELUM", "-", "Fitur belum ada."),
    ("B. Public API", "/api/public/weather/weekly-temperature", "GET", "lat, long", "x-public-token",
     "Grafik 'Suhu Mingguan'", "ForecastTemperatureChart.vue", "BELUM", "src/data/mockData.ts",
     "Respons: {data:{daily:[{date,temperature}],weekly:[...]}}."),
    ("B. Public API", "/api/public/weather/video-latest", "GET", "hashtag=infobmkgpws", "x-public-token",
     "Video Instagram terbaru BMKG", "NewsSection.vue", "BELUM", "src/data/mockData.ts (newsArticles)",
     "Respons: array {title, videoUrl, ...}."),
    ("B. Public API", "/api/public/banners", "GET", "-", "x-public-token",
     "Banner promo di dashboard", "-", "BELUM", "-", "Respons: {data:[{banner}]}. Fitur belum ada."),
    ("A. ISDP v1", "/api/v1/setting/find-code", "GET", "code=active", "Terbuka",
     "Tema background per section (public/aviation/maritim)", "src/data/cityThemes.ts", "BEDA-SUMBER",
     "Tema statis lokal", "Bebas auth; nilai setting.background bisa dipakai untuk dinamisasi tema."),
    ("A. ISDP v1", "/api/v1/sunset/json", "GET", "lat, lng", "X-API-KEY",
     "Info matahari & bulan (30+ field: sunrise, golden_hour, moon_phase, dll)",
     "CurrentWeather.vue (sunrise/sunset untuk day/night)", "BELUM", "Data statis/mock",
     "Bisa replace dengan perhitungan lokal (mis. SunCalc) ATAU langsung endpoint ini."),
    ("A. ISDP v1", "/api/v1/tcwc/cyclone/all", "GET", "-", "X-API-KEY",
     "Sirkulasi siklon TCWC Jakarta", "-", "BELUM", "-", "Respons: {status, data:{cyclone_info:[...]}}."),
    ("A. ISDP v1", "/api/v1/public/maritim/nearest-location", "GET", "lat, long", "x-public-token",
     "Wilayah laut terdekat (drawer maritim)", "WaveRadarMap.vue / MaritimeAdvisorDrawer.vue", "BELUM",
     "-", "Respons: {data:{code:'F.09', name:'Teluk Jakarta', wilpel, geometry:Polygon}}."),
    ("A. ISDP v1", "/api/v1/maritim/route", "GET", "-", "X-API-KEY",
     "Rute pelayaran maritim", "TransportWeather.vue", "BEDA-SUMBER", "OSRM (router.project-osrm.org)",
     "404 via Node, 200 via browser (butuh context CF)."),
    ("A. ISDP v1", "/api/v1/maritim/water-area/water_area_point.geojson", "GET", "-", "X-API-KEY",
     "GeoJSON titik area laut", "-", "BELUM", "-", "-"),
    ("A. ISDP v1", "/api/v1/sus/modelrun", "GET", "-", "X-API-KEY",
     "Katalog run model (radar/windmap)", "SatelliteMap.vue", "BEDA-SUMBER",
     "inderaja.bmkg.go.id/IMAGE/HIMA/*.png (statis)", "Original pakai ini + spartan untuk layer dinamis."),
    ("A. ISDP v1", "/api/v1/api/signature/impact/public/list/<YYYY-MM-DDT00:00:00Z>", "GET", "path: tanggal (7 hari)", "X-API-KEY",
     "Impact-Based Forecast per tanggal (/ibf + layer Peringatan /map)", "-", "BELUM", "-",
     "Fitur IBF belum ada di redesign."),
    ("A. ISDP v1", "/api/v1/maps/{general,warning,maritim,aviation,tourism}/metadata/tiles", "GET", "-", "campuran",
     "Metadata layer peta (sources, layers, overlays)", "SatelliteMap.vue, WaveRadarMap.vue", "BEDA-SUMBER",
     "CartoCDN basemap (basemaps.cartocdn.com)", "warning/maritim terbuka; general/aviation/tourism perlu X-API-KEY."),
    ("A. ISDP v1", "/api/v1/tourism/tiles/data/tourism.json", "GET", "-", "X-API-KEY",
     "POI wisata (layer Turis /map)", "LandBasedActivities.vue", "BEDA-SUMBER", "Overpass API (OSM)",
     "-"),
    ("A. ISDP v1", "/api/v1/sus/tiles/data/suspoi.json", "GET", "-", "X-API-KEY",
     "POI layer umum /map", "LandBasedActivities.vue", "BEDA-SUMBER", "Overpass API (OSM)",
     "502 flaky saat probe; 200 via browser."),
    ("A. ISDP v1", "/api/v1/spartan/pollution/{pol}/{lvl}/{t0}/{t1}/{z}/{x}/{y}.png", "GET", "tile z/x/y", "-",
     "Tile peta polusi (layer Peringatan /map)", "ForecastAirQualityChart.vue", "BELUM",
     "src/data/mockData.ts", "AQ chart redesign pakai data mock."),
    ("A. ISDP v1", "/api/v1/user/subscribe-type", "GET", "-", "x-public-token",
     "Tipe subscription (/premium)", "-", "BELUM", "-", "Halaman premium belum ada di redesign."),
    ("A. ISDP v1", "/api/v1/developer/{login,register,token-management,token-management/services}", "POST/GET", "email, password", "Akun dev",
     "Console developer (/devtools, /Login, /register)", "-", "BELUM", "-",
     "Tidak relevan untuk redesign publik (opsional)."),
    ("E. Eksternal Original", "spartan.bmkg.go.id/map/modelrun", "GET", "-", "Terbuka",
     "Katalog model (gfs_indo, spartan, pcm_pm25, QPE, FDRS, dll)", "-", "BELUM", "-",
     "Respons kaya (14+ katalog model) — potensi untuk layer radar/iklim."),
    ("E. Eksternal Original", "widis.bmkg.go.id/ndf/cgms/weather/forward", "GET", "params dinamis", "-",
     "Data satelit CGMS", "SatelliteMap.vue", "BEDA-SUMBER", "inderaja.bmkg.go.id/IMAGE/HIMA/*.png",
     "404 tanpa param; butuh reverse-engineering param dari browser."),
    ("E. Eksternal Original", "bmkg-sus.geo.id/windmap/ecmwf/1000", "GET", "tile", "-",
     "Tile angin ECMWF", "-", "BELUM", "-", "Wind layer belum ada di redesign."),
    ("E. Eksternal Original", "tiles.circlegeo.com/data/{indocg,administration}.json", "GET", "-", "Terbuka",
     "Boundary administratif peta", "SatelliteMap.vue", "BEDA-SUMBER", "CartoCDN basemap",
     "-"),
    ("E. Eksternal Original", "cuaca.bmkg.go.id/blog/wp-json/wp/v2/", "GET", "wp standard", "Terbuka",
     "Konten /berita, /berita/infografis, /berita/video (SSR)", "NewsSection.vue", "BELUM",
     "src/data/mockData.ts (newsArticles)", "WP REST API standar — mudah diintegrasikan."),
    ("R. Redesign-Only", "data.bmkg.go.id/DataMKG/TEWS/gempadirasakan.json", "GET", "-", "Terbuka",
     "TIDAK ADA di original (cuaca.bmkg.go.id tidak menampilkan gempa)", "EarthquakeActivity.vue, EarthquakeHistory.vue",
     "REDESIGN-ONLY", "-", "Legacy BMKG endpoint; dipakai redesign sebagai fitur tambahan."),
    ("R. Redesign-Only", "nominatim.openstreetmap.org (search/reverse)", "GET", "lat/lon/q", "-",
     "TIDAK ADA di original", "App.vue, LandBasedActivities.vue", "REDESIGN-ONLY", "-",
     "Pengganti adm/coord + pencarian lokasi."),
    ("R. Redesign-Only", "overpass-api.de/api/interpreter", "GET/POST", "BBox OSM", "-",
     "TIDAK ADA di original", "LandBasedActivities.vue", "REDESIGN-ONLY", "-",
     "POI aktivitas darat."),
    ("R. Redesign-Only", "router.project-osrm.org/route/v1/driving", "GET", "coords", "-",
     "TIDAK ADA di original", "TransportWeather.vue", "REDESIGN-ONLY", "-",
     "Routing transportasi darat."),
]

STATUS_COLOR = {"BELUM": "FFC7CE", "BEDA-SUMBER": "FFEB9C", "REDESIGN-ONLY": "BDD7EE"}
STATUS_COLOR_FONT = {"BELUM": "9C0006", "BEDA-SUMBER": "9C6500", "REDESIGN-ONLY": "1F4E79"}

HEADERS = ["No", "Family", "Endpoint", "Method", "Params", "Auth",
           "Dipakai di Original (baseline)", "Komponen Redesign", "Status Redesign",
           "Sumber Pengganti (Redesign)", "Catatan / Response Shape"]
WIDTHS = [5, 20, 52, 10, 24, 16, 44, 40, 16, 38, 60]

wb = Workbook()
ws = wb.active
ws.title = "API Mapping"

header_fill = PatternFill("solid", fgColor="1F4E79")
header_font = Font(color="FFFFFF", bold=True, size=11)
wrap = Alignment(wrap_text=True, vertical="top")

for c, (h, w) in enumerate(zip(HEADERS, WIDTHS), 1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.fill, cell.font, cell.alignment = header_fill, header_font, wrap
    ws.column_dimensions[get_column_letter(c)].width = w

for i, row in enumerate(ROWS, start=2):
    status = row[7]
    values = (i - 1,) + row
    for c, v in enumerate(values, 1):
        cell = ws.cell(row=i, column=c, value=v)
        cell.alignment = wrap
        if c == 9:  # status
            cell.fill = PatternFill("solid", fgColor=STATUS_COLOR[status])
            cell.font = Font(color=STATUS_COLOR_FONT[status], bold=True)

ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:K{len(ROWS) + 1}"

# Sheet ringkasan
ws2 = wb.create_sheet("Ringkasan")
ws2.append(["Ringkasan Komparasi API — baseline cuaca.bmkg.go.id (8 Sep 2026) vs project bmkg_cuaca"])
ws2["A1"].font = Font(bold=True, size=13)
ws2.append([])
ws2.append(["Status", "Jumlah", "Arti"])
ws2["A4"].font = ws2["B4"].font = ws2["C4"].font = Font(bold=True)
meaning = {
    "BELUM": "Endpoint original belum diimplementasikan (data masih mock / fitur belum ada)",
    "BEDA-SUMBER": "Fitur ada di redesign tapi mengambil data dari sumber berbeda (OSM/OSRM/inderaja/Carto/statik)",
    "REDESIGN-ONLY": "Sumber yang hanya dipakai redesign (tidak dipakai original)",
}
for s in ["BELUM", "BEDA-SUMBER", "REDESIGN-ONLY"]:
    ws2.append([s, sum(1 for r in ROWS if r[7] == s), meaning[s]])
    ws2.cell(row=ws2.max_row, column=1).fill = PatternFill("solid", fgColor=STATUS_COLOR[s])
ws2.append([])
ws2.append(["Total endpoint terpetakan", len(ROWS)])
ws2.append(["Endpoint original terimplementasi identik (pakai API yang sama)", 0])
ws2["A%d" % ws2.max_row].font = Font(bold=True, color="9C0006")
ws2.column_dimensions["A"].width = 50
ws2.column_dimensions["B"].width = 10
ws2.column_dimensions["C"].width = 80
for r in range(5, ws2.max_row + 1):
    ws2.cell(row=r, column=3).alignment = wrap

wb.save("api-mapping.xlsx")
print(f"OK: api-mapping.xlsx ({len(ROWS)} baris)")
print({s: sum(1 for r in ROWS if r[7] == s) for s in STATUS_COLOR})
